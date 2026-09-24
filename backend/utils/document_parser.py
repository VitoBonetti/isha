import io
import zipfile
import pandas as pd
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
import google.auth
from audit_logger import log_audit_event

# --- MIME TYPE MAPPINGS ---
GOOGLE_MIME_TYPES = {
    'doc': 'application/vnd.google-apps.document',
    'sheet': 'application/vnd.google-apps.spreadsheet',
    'slide': 'application/vnd.google-apps.presentation'
}

EXPORT_MIME_TYPES = {
    'doc': 'text/plain',
    'sheet': 'text/csv',
    'slide': 'text/plain'
}


def get_drive_service():
    # Use your existing GCP service account credentials setup here
    credentials, _ = google.auth.default(scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build('drive', 'v3', credentials=credentials)


def extract_text_from_drive_file(drive_file_id: str, mime_type: str) -> str:
    """
    Downloads or exports a file from Google Drive and extracts all text.
    Handles both native Google Workspace files (Google Slides/Docs/Sheets)
    and standard binary uploads (.pptx, .pdf, .docx).
    """
    service = get_drive_service()
    real_mime_type = mime_type

    # 1. FETCH LIVE METADATA FROM GOOGLE DRIVE (Size + Real MIME Type)
    try:
        file_metadata = service.files().get(fileId=drive_file_id, fields="size, mimeType").execute()
        file_size_bytes = int(file_metadata.get('size', 0))

        # Override DB mime_type with the actual MIME type from Google Drive
        if file_metadata.get('mimeType'):
            real_mime_type = file_metadata.get('mimeType')

        MAX_FILE_SIZE_MB = 25
        if file_size_bytes > (MAX_FILE_SIZE_MB * 1024 * 1024):
            log_audit_event(
                user_id="SYSTEM", role="SYSTEM", action="DOCUMENT_PARSING_EXTRACT_TEXT_ERROR",
                resource_type="DOCUMENT_PARSING", resource_id="DOCUMENT_PARSING",
                details=f"File exceeds maximum allowed size of {MAX_FILE_SIZE_MB}MB."
            )
            raise ValueError(f"File exceeds maximum allowed ingestion size of {MAX_FILE_SIZE_MB}MB.")
    except ValueError as ve:
        raise ve
    except Exception as ex:
        log_audit_event(
            user_id="SYSTEM", role="SYSTEM", action="DOCUMENT_PARSING_EXTRACT_EXCEPTION",
            resource_type="DOCUMENT_PARSING", resource_id="DOCUMENT_PARSING",
            details=f"Ignore standard API fetch errors and proceed to download attempt: {ex}"
        )

    # 2. HANDLE NATIVE GOOGLE WORKSPACE FILES (EXPORT METHOD)
    if real_mime_type == GOOGLE_MIME_TYPES['doc']:
        request = service.files().export_media(fileId=drive_file_id, mimeType=EXPORT_MIME_TYPES['doc'])
        return request.execute().decode('utf-8', errors='ignore')

    elif real_mime_type == GOOGLE_MIME_TYPES['sheet']:
        request = service.files().export_media(fileId=drive_file_id, mimeType=EXPORT_MIME_TYPES['sheet'])
        return request.execute().decode('utf-8', errors='ignore')

    elif real_mime_type == GOOGLE_MIME_TYPES['slide']:
        request = service.files().export_media(fileId=drive_file_id, mimeType=EXPORT_MIME_TYPES['slide'])
        return request.execute().decode('utf-8', errors='ignore')

    # 3. HANDLE STANDARD BLOB FILES (DOWNLOAD METHOD WITH FALLBACK)
    try:
        request = service.files().get_media(fileId=drive_file_id)
        file_bytes = request.execute()
    except HttpError as err:
        # Fallback: If get_media fails because the file is non-binary, attempt plain text export
        if "fileNotDownloadable" in str(err) or err.resp.status == 403:
            request = service.files().export_media(fileId=drive_file_id, mimeType='text/plain')
            return request.execute().decode('utf-8', errors='ignore')
        raise err

    file_stream = io.BytesIO(file_bytes)

    # 4. PARSE BLOB FILES BASED ON REAL MIME TYPE
    if real_mime_type == 'application/pdf':
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        return "\n".join([page.get_text() for page in doc])

    elif real_mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':  # .docx
        doc = Document(file_stream)
        return "\n".join([para.text for para in doc.paragraphs])

    elif real_mime_type == 'application/vnd.openxmlformats-officedocument.presentationml.presentation':  # .pptx
        prs = Presentation(file_stream)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            text_runs.append(paragraph.text.strip())
        return "\n".join(text_runs)

    elif real_mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':  # .xlsx
        df = pd.read_excel(file_stream)
        return df.to_csv(index=False)

    elif real_mime_type == 'application/zip':
        return parse_zip_file(file_stream)

    elif real_mime_type.startswith('text/') or real_mime_type in [
        'application/json', 'application/javascript', 'application/csv',
        'text/csv', 'text/markdown', 'text/x-markdown'
    ]:
        return file_bytes.decode('utf-8', errors='ignore')

    else:
        return ""


def parse_zip_file(zip_stream: io.BytesIO) -> str:
    """Extracts text from all readable code/text files inside a ZIP archive."""
    text_content = []
    ignore_dirs = ['node_modules/', '.git/', '__pycache__/']
    valid_extensions = ('.py', '.js', '.ts', '.java', '.json', '.txt', '.md', '.html', '.css', '.go', '.rs')

    with zipfile.ZipFile(zip_stream) as z:
        for filename in z.namelist():
            if filename.endswith('/') or any(ign in filename for ign in ignore_dirs):
                continue
            if filename.endswith(valid_extensions):
                try:
                    with z.open(filename) as f:
                        file_text = f.read().decode('utf-8', errors='ignore')
                        text_content.append(f"--- FILE: {filename} ---\n{file_text}\n")
                except Exception as ex:
                    log_audit_event(
                        user_id="SYSTEM",
                        role="SYSTEM",
                        action="DOCUMENT_PARSING_ZIP_FILE_EXCEPTION",
                        resource_type="DOCUMENT_PARSING",
                        resource_id="PARSING_ZIP",
                        details=f"Exception parsing zip file: {ex}. Continue"
                    )
                    continue

    return "\n".join(text_content)