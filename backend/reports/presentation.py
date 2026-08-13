import requests, datetime, os, re, unicodedata
from datetime import date, datetime, timezone, timedelta
from io import BytesIO
from typing import List, Dict, Any, Union
from collections import Counter
import urllib3
from urllib3.exceptions import InsecureRequestWarning
from bs4 import BeautifulSoup
from PIL import Image
import html2text
import google.auth
import json
import os
from pptx import Presentation
from pptx.util import Inches, Cm, Pt
from pptx.chart.data import CategoryChartData, ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT, MSO_AUTO_SIZE
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from google.cloud import secretmanager
import google.generativeai  as genai
from  utils.secret_manager import get_secret


TEMPLATE_FOLDER_ID = os.environ.get("PRESENTATION_TEMPLATE_FOLDER_ID")
TEMPLATE_PPTX_NAME = os.environ.get("TEMPLATE_PPTX_NAME")
VULN_CYCLE_IMAGE_NAME = os.environ.get("VULN_CYCLE_IMAGE_NAME")
KEEPITSECURE_URL = os.environ.get("KISS_24_ENDPOINT")
KISS_KEY = os.environ.get("KISS_24_API_KEY_NAME")
LUIGI_KEY_NAME = os.environ.get("LUIGI_KEY_NAME")

# Define a dictionary of replacements
replacements = {
    "[TARGET_NAME]": "",
    "[ASSET_ID]": "",
    "[SERVICE_TYPE]": "",
    "[TIMESTAMP]": f"Generated on {datetime.now().strftime('%d-%m-%Y')}",
    "[OPCO]": "",
    "[DURATION]": "",
    "[START_DATE]": "DD-MM-YYYY",
    "[END_DATE]": "DD-MM-YYYY",
    "[TOTAL_VULNS]": "",
    "[VULNS_SUMMARY]": "",
    "[VULN_STATUS]": "",
    "[VULN_DUE]": "DD-MM-YYYY",
    "[REQUEST_ID]": "",
    "[MANAGEMENT_SUMMARY]": "",
}

# Service types
services = {
    "7ea682fb-e03e-43c8-a476-877d088b6cdb": "Black Box",
    "ff52df96-1934-4cc0-9090-804dc9feeae9": "Grey Box",
    "e14f6754-8b1f-41e7-88dd-cbb4fda7a3f2": "White Box",
    "43b55a12-bfb3-4e2d-a212-0db9d73358e4": "Adversary Simulation",
}

group_of_tests = {
    "#": ("Group", "E.g"),
    "1": ("Information Gathering", "Fingerprint Web Server, Fingerprint Web Application, Map Application Architecture"),
    "2": ("Configuration and Deploy Management Testing", "Application Platform Configuration, HSTS, HTTP Methods, Backup and Unreferenced Files"),
    "3": ("Identity Management Testing", "User Registration Process, Account Enumeration"),
    "4": ("Authentication Testing", "Default Credentials, Weak Lock Out Mechanism, Reset Password, Weak Password Policy"),
    "5": ("Authorization Testing", "Directory Traversal/File Inclusion, Insecure Direct Object References, Privilege Escalation"),
    "6": ("Session Management Testing", "Session Fixation, Logout Functionality, Session Timeout, Bypassing Session Management Schema"),
    "7": ("Data Validation Testing", "SQL Injection, Code Injection, Reflected XSS, Stored XSS, HTTP Parameter Pollution"),
    "8": ("Error Handling", "Analysis of Error Codes, Analysis of Stack Traces"),
    "9": ("Cryptography", "Weak SSL/TLS Ciphers, Sensitive information sent via unencrypted channels"),
    "10": ("Business Logic Testing", "Ability to Forge Requests, Business Logic Data Validation, Upload of Malicious Files"),
    "11": ("Client Side Testing", "JavaScript Execution, HTML Injection, Client Side URL Redirect, Clickjacking")
}

## Colors
gColors = {
    'darkred': RGBColor(168, 0, 0),
    'red': RGBColor(231, 69, 54),
    'orange': RGBColor(255, 181, 17),
    'green': RGBColor(0, 204, 102),
    'yellow': RGBColor(255, 253, 173),
    'white': RGBColor(255, 255, 255),
    'black': RGBColor(0, 0, 0),
    'blue': RGBColor(33, 117, 217),
}

gColorsRisk = {
    'Critical': gColors['darkred'],
    'High': gColors['red'],
    'Medium': gColors['orange'],
    'Low': gColors['yellow'],
    'Info': gColors['blue']
}


# --- Google Gemini function for management summary creation --
def extract_sections(html: str) -> Dict[str, str]:
    """
    Pull out Description, Impact and Recommendation text from an HTML blob.
    """
    soup = BeautifulSoup(html, 'html.parser')
    sections: Dict[str, str] = {}

    for name in ("Description", "Impact", "Recommendation"):
        heading = soup.find("strong", string=name)
        if not heading:
            continue

        h1 = heading.find_parent("h1")
        if not h1:
            continue

        parts = []
        for sib in h1.next_siblings:
            if getattr(sib, "name", None) == "h1":
                break
            parts.append(str(sib))

        text = BeautifulSoup("".join(parts), "html.parser").get_text(separator=" ", strip=True)
        sections[name] = text

    return sections


def fetch_vulns_details(session, headers, testUUID) -> List[Dict[str, Union[str, Any]]]:
    try:
        resp = session.post(f"{KEEPITSECURE_URL}vulnerabilities", json={"tests": [testUUID]}, headers=headers, verify=False)
        resp.raise_for_status()
    except requests.RequestException as e:
        print(f"Failed to fetch vuln details for {testUUID}: {e}")
        return []

    items = resp.json().get("items", [])
    cleaned = []
    for item in items:
        details_html = item.get("details", "")
        sections = extract_sections(details_html)
        cleaned.append({
            "title": item.get("description", ""),
            "severity": item.get("severity", ""),
            **sections
        })

    return cleaned


def build_management_summary_payload(
    vulns: List[Dict[str, Any]]
) -> Dict[str, Any]:
    """
    Build a summary dict: total count + ordered severity counts.
    """
    order = ["Critical", "High", "Medium", "Low", "Info"]
    counts = Counter(v["severity"] for v in vulns)

    return {
        "total": len(vulns),
        "severity_counts": [counts.get(level, 0) for level in order],
        "order": order,
        "vulnerabilities": vulns,
    }


def ai_generate_management_summary(data):
    api_key = get_secret(LUIGI_KEY_NAME)
    genai.configure(api_key=api_key)
    prompt = """You will receive a JSON object containing information about vulnerabilities discovered during a security assessment. Your task is to generate a concise management summary suitable for executives. 

        The summary must follow this structure:

        1. Start with:  
           "During our most recent assessment we have identified <TOTAL NUMBER OF FINDINGS> issue(s): <NUMBER OF FINDINGS PER SEVERITY>"  
           Use the 'order' and 'severity_counts' fields in the JSON to produce the correct severity breakdown.  Ignore the one where the 'severity_counts' is 0 (zero).

        2. After that, write a **single-sentence** management summary that describes the main security concerns and their potential impact, based on all the "vulnerabilities" objects provided. The summary must be fluent, natural, and not formatted as a list or use any markdown. Use professional but clear language.

        3. Conclude the summary with:  
           "Addressing these issues will result in a better security posture."

        Return your result as a JSON object like this:  
        {"summary": "<GENERATED SUMMARY>"}
        Do not wrap the output in any markdown, code blocks, or additional text. Return only the plain JSON object."""

    model = genai.GenerativeModel(
        'gemini-3.1-flash-lite',
        system_instruction=prompt
    )

    response = model.generate_content(
        json.dumps(data),
        generation_config=genai.types.GenerationConfig(temperature=0.7)
    )
    ai_response = response.text.strip()

    # Try to parse as JSON and extract 'summary'; fallback to raw string if not JSON
    try:
        ai_response_dict = json.loads(ai_response)
        summary = ai_response_dict.get("summary", ai_response)  # Use 'summary' if present, else full response
    except json.JSONDecodeError:
        summary = ai_response  # Not JSON, so use the raw output

    print(summary)  # Keep the print if you want it for logging/debugging
    return summary or ""  # Return the summary string (or empty if somehow None/empty)


# --- Health Check and Update information ---
# Function that fetch the test info THIS IS A DUPLICATE FUNCTION CAN BE EVENTUALLY WITH A SMALL MODIFYCATION BE DELETED
def fetch_test_info(session, headers, testUUID):
    try:
        resp = session.post(f"{KEEPITSECURE_URL}tests", json={"uuid": [testUUID]}, headers=headers,
                            verify=False)
        resp.raise_for_status()
        items = resp.json().get("items", [])

        return json.dumps(items, indent=4)

    except requests.exceptions.HTTPError as e:
        print(e)
    except requests.exceptions.RequestException as e:
        print(e)
    except Exception as e:
        print(e)


def end_test(session, headers, testUUID):
    try:
        resp = session.post(f"{KEEPITSECURE_URL}tests/{testUUID}/change-state", json={"state": "end"}, headers=headers,
                            verify=False)
        resp.raise_for_status()
    except requests.exceptions.HTTPError as e:
        print(e)
    except requests.exceptions.RequestException as e:
        print(e)
    except Exception as e:
        print(e)


# Function that check the health of the test
def check_test_health(session, headers, testUUID):
    object = json.loads(fetch_test_info(session, headers, testUUID))

    # Dictionary for Healthy check
    not_healthy = {}

    # get interested fields
    try:
        state = object[0].get("state")
        # check state health
        not_healthy["state"] = {}
        if state != "Tested":
            end_test(session, headers, testUUID)
            state = "Tested"
            not_healthy["state"]["healthy"] = False
            not_healthy["state"]["reason"] = f"{state}"
        else:
            not_healthy["state"]["healthy"] = True

        details = object[0].get("details")
        started_at = object[0].get("started_at")
        ended_at = object[0].get("ended_at")
    except Exception as e:
        print("[-] It's not possible to generate a slide  from this test. Please check your test and try again.")
        exit(1)

    # check started_at health
    not_healthy["started_at"] = {}
    try:
        datetime.strptime(started_at, "%Y-%m-%d %H:%M:%S")
        not_healthy["started_at"]["healthy"] = True
    except:
        not_healthy["started_at"]["healthy"] = False
        not_healthy["started_at"]["reason"] = f"Unhealthy started_at: {started_at}"

    # check ended_at health
    not_healthy["ended_at"] = {}
    try:
        datetime.strptime(ended_at, "%Y-%m-%d %H:%M:%S")
        not_healthy["ended_at"]["healthy"] = True
    except:
        not_healthy["ended_at"]["healthy"] = False
        not_healthy["ended_at"]["reason"] = f"Unhealthy ended_at: {ended_at}"

    try:
        soup = BeautifulSoup(details, "html.parser")
        # extracts all the text, stripping away the HTML tags
        text = soup.get_text()

        # Remove unwanted carriage returns, newlines, tabs, and non-breaking spaces
        # Then, normalize multiple spaces into a single space for clean output
        cleaned_text = text.replace('\r', '').replace('\n', ' ').replace('\t', '')
        cleaned_text = ' '.join(cleaned_text.split())

        # Check if "[DETAILS]" is in the cleaned text
        not_healthy["details"] = {}
        if "[DETAILS]" in cleaned_text:
            not_healthy["details"]["healthy"] = True
        else:
            not_healthy["details"]["healthy"] = False
            not_healthy["details"]["reason"] = "No [DETAILS] marker"

        # Check if "[MANAGEMENT SUMMARY]" is in the cleaned text
        not_healthy["summary"] = {}
        if "[MANAGEMENT SUMMARY]" in cleaned_text:
            # Check if there is text after the marker "[MANAGEMENT SUMMARY]"
            marker = "[MANAGEMENT SUMMARY]"
            text_after_marker = cleaned_text.split(marker, 1)[1]
            if text_after_marker.strip():
                not_healthy["summary"]["healthy"] = True
            else:
                not_healthy["summary"]["healthy"] = False
                not_healthy["summary"]["reason"] = "Not Management Summary"
        else:
            not_healthy["summary"]["healthy"] = False
            not_healthy["summary"]["reason"] = "Not [MANAGEMENT SUMMARY] Marker. Summary has been generated with AI."

        # Check if "ServiceNow Request ID" is in the cleaned text
        not_healthy["service-now"] = {}
        if "ServiceNow Request ID:" in cleaned_text:
            not_healthy["service-now"]["healthy"] = True
        else:
            not_healthy["service-now"]["healthy"] = False
            not_healthy["service-now"]["reason"] = "No Service Now Request ID marker"

        # Check if "Connection Type" is in the cleaned text
        not_healthy["connection-type"] = {}
        if "Connection Type:" in cleaned_text:
            not_healthy["connection-type"]["healthy"] = True
        else:
            not_healthy["connection-type"]["healthy"] = False
            not_healthy["connection-type"]["reason"] = "No Connection Type marker"

        # Check if "Test Accounts" is in the cleaned text
        not_healthy["accounts"] = {}
        if "Test Accounts:" in cleaned_text:
            not_healthy["accounts"]["healthy"] = True
        else:
            not_healthy["accounts"]["healthy"] = False
            not_healthy["accounts"]["reason"] = "No Test Accounts marker"

        # Check if "Account Roles" is in the cleaned text
        not_healthy["roles"] = {}
        if "Account Roles:" in cleaned_text:
            not_healthy["roles"]["healthy"] = True
        else:
            not_healthy["roles"]["healthy"] = False
            not_healthy["roles"]["reason"] = "No Account Roles marker"
    except Exception as ex:
        print(ex)
        not_healthy["details"]["healthy"] = False
        not_healthy["details"]["reason"] = "No details"

    return not_healthy, cleaned_text


def analyze_and_fix(healthy_check, details, session, headers, testUUID):
    data_info = json.dumps(healthy_check, indent=4)
    data = json.loads(data_info)
    warning_list = []
    test_details_payload = ""

    if not data["state"]["healthy"]:
        if data["state"]["reason"] == "New":
            print("[-] Test not started yet. Exiting")
            exit(1)
        else:
            warning_list.append(f"State set as {data['state']['reason']}.")

    if not data["started_at"]["healthy"]:
        warning_list.append(f"Started_at set as {data['started_at']['reason']}.")

    if not data["ended_at"]["healthy"]:
        warning_list.append(f"Ended_at set as {data['ended_at']['reason']}.")

    if not data["details"]["healthy"]:
        warning_list.append(f"Details set as {data['details']['reason']}.")

    test_details_payload = "<p>[DETAILS]<br><br>"


    if not data["service-now"]["healthy"]:
        warning_list.append(f"service-now set as {data['service-now']['reason']}.")
        test_details_payload += "ServiceNow Request ID: N/A<br>"
    else:
        match = re.search(r'ServiceNow Request ID: (.*?) Connection Type:', details)
        if match:
            ritm_id = match.group(1).strip()
            test_details_payload += f"ServiceNow Request ID: {ritm_id}<br>"
        else:
            test_details_payload += "ServiceNow Request ID: N/A<br>"

    if not data["connection-type"]["healthy"]:
        warning_list.append(f"connection-type set as {data['connection-type']['reason']}.")
        test_details_payload += "Connection Type: N/A<br>"
    else:
        match = re.search(r'Connection Type: (.*?) Test Accounts:', details)
        if match:
            ritm_id = match.group(1).strip()
            test_details_payload += f"Connection Type: {ritm_id}<br>"
        else:
            test_details_payload += "Connection Type: N/A<br>"

    if not data["accounts"]["healthy"]:
        warning_list.append(f"service-now set as {data['accounts']['reason']}.")
        test_details_payload += "Test Accounts: N/A<br>"
    else:
        match = re.search(r'Test Accounts: (.*?) Account Roles:', details)
        if match:
            ritm_id = match.group(1).strip()
            test_details_payload += f"Test Accounts: {ritm_id}<br>"
        else:
            test_details_payload += "Test Accounts: N/A<br>"

    if not data["roles"]["healthy"]:
        warning_list.append(f"roles set as {data['roles']['reason']}.")
        test_details_payload += "Account Roles: N/A<br></p>"
    else:
        match = re.search(r'Account Roles: (.*?) \[MANAGEMENT SUMMARY]', details)
        if match:
            ritm_id = match.group(1).strip()
            test_details_payload += f"Account Roles: {ritm_id}<br></p>"
        else:
            test_details_payload += "Account Roles: N/A<br></p>"

    if not data["summary"]["healthy"]:
        warning_list.append(f"Summary set as {data['summary']['reason']}.")
        test_details_payload += "<p>[MANAGEMENT SUMMARY] <br><br>"

        data = build_management_summary_payload(fetch_vulns_details(session, headers, testUUID))
        ai_payload = ai_generate_management_summary(data)
        test_details_payload += f"{ai_payload}</p>"
    else:
        match = re.search(r'\[MANAGEMENT SUMMARY] (.*)', details)
        if match:
            ritm_id = match.group(1).strip()
            test_details_payload += "<p>[MANAGEMENT SUMMARY] <br><br>"
            test_details_payload += f"{ritm_id}<br></p>"
    if warning_list:
        response = session.patch(f"{KEEPITSECURE_URL}tests/{testUUID}/edit", headers=headers,
                                json={"details": test_details_payload}, verify=False)
        response.raise_for_status()
        if response.status_code == 200:
            print("[+] Test details updated successfully.")
            return warning_list
        else:
            print("[-] Test details update failed.")
            print(f"[!] Test details update failed with status code: {response.status_code}.")
            print(f"[!] Test details update failed with details: {response.text}.")
            exit(1)
    else:
        print("[+] Test is healthy, no edit required.")
        return []


# --- Google Cloud Helper Functions ---
def get_drive_service():
    """Authenticates using default service account credentials and returns a Google Drive service object."""
    credentials, project = google.auth.default(scopes=['https://www.googleapis.com/auth/drive'])
    return build('drive', 'v3', credentials=credentials)


def get_file_id_by_name(service, folder_id, file_name):
    """Finds a file by name in a specific Google Drive folder and returns its ID."""
    try:
        query = f"'{folder_id}' in parents and name = '{file_name}' and trashed = false"
        response = service.files().list(q=query, spaces='drive', fields='files(id, name)', supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        files = response.get('files', [])
        if not files:
            raise FileNotFoundError(f"File '{file_name}' not found in Drive folder ID '{folder_id}'.")
        return files[0].get('id')
    except HttpError as error:
        print(f"An error occurred while searching for file '{file_name}': {error}")
        raise


def download_drive_file(service, file_id):
    """Downloads a file from Google Drive and returns it as a BytesIO object."""
    try:
        request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
        file_stream = BytesIO()
        downloader = MediaIoBaseDownload(file_stream, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
        file_stream.seek(0)
        return file_stream
    except HttpError as error:
        print(f"An error occurred while downloading file ID '{file_id}': {error}")
        raise


def upload_drive_file(service, folder_id, filename, file_stream):
    """Uploads a file stream to a specific Google Drive folder."""
    try:
        file_metadata = {'name': filename, 'parents': [folder_id]}
        media = MediaIoBaseUpload(
            file_stream,
            mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
            resumable=True
        )
        file = service.files().create(body=file_metadata, media_body=media, fields='id, webViewLink', supportsAllDrives=True).execute()
        return file.get('webViewLink')
    except HttpError as error:
        print(f"An error occurred while uploading file '{filename}': {error}")
        raise


# --- Presentation & API Helper Functions ---
def _add_slide(prs, layout_name, slide_layouts):
    """Adds a new slide to the presentation using a named layout."""
    layout_index = slide_layouts.get(layout_name)
    if layout_index is None:
        print(f"Error: Slide layout '{layout_name}' not found in template.")
        raise ValueError(f"Slide layout '{layout_name}' not found.")
    return prs.slides.add_slide(prs.slide_layouts[layout_index])


def retrieve_test_info(session, headers, testUUID):
    resp = session.post(f"{KEEPITSECURE_URL}tests", json={"uuid": [testUUID]}, headers=headers, verify=False)
    resp.raise_for_status()
    return resp.json()["items"][0]


def retrieve_test_service(session, headers, testUUID):
    resp = session.post(f"{KEEPITSECURE_URL}tags", json={"tests": testUUID}, headers=headers, verify=False)
    resp.raise_for_status()
    for item in resp.json()["items"]:
        if item["uuid"] in services:
            return services[item["uuid"]]
    print("[!] No service type found.")
    return ""


def retrieve_asset_info(session, headers, testUUID):
    resp = session.post(f"{KEEPITSECURE_URL}assets", json={"tests": testUUID}, headers=headers, verify=False)
    resp.raise_for_status()
    return resp.json()["items"]


def retrieve_onetrust_id(session, headers, assetUUID):
    resp = session.post(f"{KEEPITSECURE_URL}fields", json={"assets": assetUUID}, headers=headers, verify=False)
    resp.raise_for_status()
    for item in resp.json()["items"]:
        if item["custom_field"]["name"] == "OneTrust ID":
            return item.get("value", "")
    return ""


def retrieve_vulns_by_test(session, headers, testUUID):
    resp = session.post(f"{KEEPITSECURE_URL}vulnerabilities", json={"tests": testUUID}, headers=headers, verify=False)
    resp.raise_for_status()
    return resp.json()["items"]


def is_image_stream(stream):
    """Checks if a byte stream is a PNG or JPG."""
    header = stream.read(4)
    stream.seek(0)  # Reset stream for later use
    return header.startswith(b'\x89\x50\x4E\x47') or header.startswith(b'\xFF\xD8')


def _add_image(presentation, slide_layouts, image_stream):
    """Adds an image from a byte stream to a new 'evidence' slide."""
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "evidence"

    im = Image.open(image_stream)
    width, height = im.size

    max_width = presentation.slide_width - Inches(2)
    max_height = presentation.slide_height - Inches(2)
    ratio = min(max_width / width, max_height / height)
    pic_width = int(width * ratio)
    pic_height = int(height * ratio)

    left = (presentation.slide_width - pic_width) / 2
    top = (presentation.slide_height - pic_height) / 2

    image_stream.seek(0)
    slide.shapes.add_picture(image_stream, left, top, width=pic_width, height=pic_height)


def retrieve_vuln_attachments(presentation, slide_layouts, session, headers, vulnUUID):
    """Retrieves and adds vulnerability attachment images to the presentation."""
    try:
        resp = session.post(f"{KEEPITSECURE_URL}attachments", json={"vulnerabilities": [vulnUUID]}, headers=headers, verify=False)
        resp.raise_for_status()
        attachments_metadata = resp.json().get("items", [])
        print(attachments_metadata)

        for attachment in attachments_metadata:
            uuid = attachment.get("uuid")
            if not uuid:
                continue
            resp_img = session.get(f"{KEEPITSECURE_URL}attachments/{uuid}", headers={"x-api-key": get_secret(KISS_KEY)}, verify=False)
            if resp_img.status_code == 200:
                image_stream = BytesIO(resp_img.content)
                if is_image_stream(image_stream):
                    _add_image(presentation, slide_layouts, image_stream)
    except Exception as error:
        print(f"[CUSTOM]: {error}")


def slugify(value):
    value = str(value)
    value = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode('ascii')
    value = re.sub(r'[^\w\s-]', '', value).strip().lower()
    return re.sub(r'[-\s]+', '-', value)


def replace_text_in_slide(slide, replacements_dict):
    """Replaces text in slide shapes with provided replacements, with debugging."""
    replaced = set()
    # Log all text in the slide for debugging
    slide_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    slide_text.append(r.text)
    print(f"Slide text content: {slide_text}")

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                for old, new in replacements_dict.items():
                    if old in r.text:
                        print(f"Replacing '{old}' with '{new}' in slide")
                        r.text = r.text.replace(old, str(new))
                        replaced.add(old)
    # Log any placeholders that were not replaced
    for placeholder in replacements_dict:
        if placeholder not in replaced:
            print(f"Warning: Placeholder '{placeholder}' not found in slide text")
    return replaced


def chart(presentation, slide_index, critical, high, medium, low, info):
    chart_data = CategoryChartData()
    chart_data.categories = ['']
    chart_data.add_series('Critical', (int(critical),))
    chart_data.add_series('High', (int(high),))
    chart_data.add_series('Medium', (int(medium),))
    chart_data.add_series('Low', (int(low),))
    chart_data.add_series('Info', (int(info),))
    x, y, cx, cy = Inches(2.7), Inches(2.5), Inches(8), Inches(4.5)
    slide = presentation.slides[slide_index]
    graphic_frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data)
    chart = graphic_frame.chart
    colors = {0: gColorsRisk['Critical'], 1: gColorsRisk['High'], 2: gColorsRisk['Medium'], 3: gColorsRisk['Low'],
              4: gColorsRisk['Info']}
    for idx, serie in enumerate(chart.plots[0].series):
        fill = serie.format.fill
        fill.solid()
        fill.fore_color.rgb = colors[idx]
    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.font.size = Pt(16)
    data_labels.font.color.rgb = RGBColor(0x0A, 0x42, 0x80)
    data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.size = Pt(13)
    chart.legend.include_in_layout = False
    plot.gap_width = 200
    vertical_axis = chart.value_axis
    vertical_axis.has_major_gridlines = False
    vertical_axis.tick_labels.font.size = Pt(10)
    vertical_axis.major_unit = 1.0


def vulns_summary_text(vulns, presentation):
    severities = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0, "Info": 0}
    for vuln in vulns:
        if vuln["severity"] in severities:
            severities[vuln["severity"]] += 1

    chart(presentation, 5, severities['Critical'], severities['High'], severities['Medium'], severities['Low'],
          severities['Info'])

    total_with_counts = {k: v for k, v in severities.items() if v > 0}
    if not total_with_counts:
        return "."

    message = ", including "
    items = list(total_with_counts.items())
    for i, (key, value) in enumerate(items):
        plural = "ies" if value > 1 else "y"
        message += f"{value} {key} Severit{plural}"
        if i < len(items) - 2:
            message += ", "
        elif i == len(items) - 2:
            message += " and "

    message += ". The total number of vulnerabilities in each category and applications has been summarized as follows:"
    return message


# --- Generate presentation ---
def generate_presentation(test_uuid: str, db_drive_folder_id: str, db_service_name: str,
                                   db_snow_number: str, start_week: int, start_year: int, duration_weeks: float) -> dict:
    """
    Main entry point for the backend to call.
    """
    print(f"Starting presentation generation for test ID: {test_uuid}")
    api_key = get_secret(KISS_KEY)

    headers = {"Content-Type": "application/json", "Accept": "application/json", "x-api-key": api_key}
    session = requests.Session()

    healthy_check, details = check_test_health(session, headers, test_uuid)
    warning_list = analyze_and_fix(healthy_check, details, session, headers, test_uuid)

    print("Connecting to Google Drive...")
    drive_service = get_drive_service()
    template_id = get_file_id_by_name(drive_service, TEMPLATE_FOLDER_ID, TEMPLATE_PPTX_NAME)
    image_id = get_file_id_by_name(drive_service, TEMPLATE_FOLDER_ID, VULN_CYCLE_IMAGE_NAME)

    print("Downloading template from Drive...")
    template_stream = download_drive_file(drive_service, template_id)
    vuln_life_cycle_image_stream = download_drive_file(drive_service, image_id)

    presentation = Presentation(template_stream)
    slide_layouts = {layout.name: i for i, layout in enumerate(presentation.slide_layouts)}

    print("Fetching data from VulnManager API...")
    test = retrieve_test_info(session, headers, test_uuid)
    asset = retrieve_asset_info(session, headers, test_uuid)[0]
    vulns = retrieve_vulns_by_test(session, headers, test_uuid)

    # Use the database values, fallback to API if DB is empty
    test_service = db_service_name if db_service_name else retrieve_test_service(session, headers, test_uuid)
    onetrust_id = db_snow_number if db_snow_number else retrieve_onetrust_id(session, headers, asset["uuid"])

    try:
        test_start = datetime.fromisocalendar(start_year, start_week, 1)
        dur_weeks = max(1, int(duration_weeks or 1))
        test_end = test_start + timedelta(days=(dur_weeks - 1) * 7 + 4)
        start_date_str = test_start.strftime("%d/%m/%Y")
        end_date_str = test_end.strftime("%d/%m/%Y")
        duration_days = int((duration_weeks or 1) * 5)
    except Exception as e:
        print(f"Warning: Could not calculate dates from planner: {e}")
        start_date_str = "N/A"
        end_date_str = "N/A"
        duration_days = 0

    replacements = {
        "[TIMESTAMP]": f"Generated on {datetime.now().strftime('%d-%m-%Y')}",
        "[DURATION]": f"{duration_days} day(s)",
        "[OPCO]": test["organisation"]["name"],
        "[SERVICE_TYPE]": test_service,
        "[TARGET_NAME]": asset["name"],
        "[ASSET_ID]": onetrust_id,
        "[START_DATE]": start_date_str,
        "[END_DATE]": end_date_str,
        "[TOTAL_VULNS]": str(len(vulns)),
        "[VULNS_SUMMARY]": vulns_summary_text(vulns, presentation)
    }

    h_parser = html2text.HTML2Text()
    h_parser.body_width = 0
    details_text = h_parser.handle(test.get('details', ''))
    req_id_match = re.search(r"ServiceNow Request ID:\s*(\S+)", details_text)
    replacements["[REQUEST_ID]"] = req_id_match.group(1) if req_id_match else "N/A"

    mgmt_summary_match = re.search(r'\[MANAGEMENT SUMMARY\]\s*\n(.*?)(?=\n#|\Z)', details_text, re.DOTALL | re.I)
    if mgmt_summary_match:
        replacements["[MANAGEMENT_SUMMARY]"] = mgmt_summary_match.group(1).strip()
    else:
        replacements["[MANAGEMENT_SUMMARY]"] = "No summary provided."

    # --- Build Presentation ---
    print("Populating presentation slides...")

    for slide in presentation.slides:
        replace_text_in_slide(slide, replacements)

    severity_order = ["critical", "high", "medium", "low", "info"]
    sorted_vulns = sorted(vulns, key=lambda v: severity_order.index(v['severity'].lower()))

    # LIST OF FINDINGS slide
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "list of findings"
    shape = slide.shapes.add_table(len(sorted_vulns) + 1, 3, Inches(0.8), Inches(1.7), Inches(11), Inches(1))
    table = shape.table
    table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(7), Inches(2), Inches(2)
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(0, 2).text = "finding", "severity", "status"
    for i, vuln in enumerate(sorted_vulns):
        table.cell(i + 1, 0).text = vuln['description']
        table.cell(i + 1, 1).text = vuln['severity']
        state = vuln['state']
        table.cell(i + 1, 2).text = "Open" if state in ["Unpublished", "Ready To Publish", "New"] else state

    # VULN DETAIL slides
    for vuln in sorted_vulns:
        slide = _add_slide(presentation, "detailed vuln", slide_layouts)

        soup = BeautifulSoup(vuln['details'], 'html.parser')
        sections = {}
        for h1 in soup.find_all('h1'):
            content = []
            for sibling in h1.find_next_siblings():
                if sibling.name == 'h1':
                    break
                content.append(sibling.get_text(strip=True))
            sections[h1.get_text(strip=True)] = "\n".join(content)

        try:
            a, b, c, d, e, f, g, h_pl = slide.placeholders
            a.text = vuln["description"] or "No title provided"
            b.text = vuln["severity"] or "Unknown"
            c.text = "impact"
            d.text = sections.get("Impact", "Not provided.")
            e.text = "recommendation"
            f.text = sections.get("Recommendation", "Not provided.")
            g.text = "description"
            h_pl.text = sections.get("Description", "Not provided.")

            # Set severity color
            fill = b.fill
            fill.solid()
            fill.fore_color.rgb = gColorsRisk.get(vuln["severity"], gColors['black'])
            if vuln["severity"] == "Low":
                b.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)

            # Set word wrap and auto size
            a.text_frame.word_wrap = True
            a.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            b.text_frame.word_wrap = True
            b.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            d.text_frame.word_wrap = True
            d.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            f.text_frame.word_wrap = True
            f.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            h_pl.text_frame.word_wrap = True
            h_pl.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except ValueError as err:
            print(f"Error accessing placeholders for vulnerability {vuln['description']}: {err}")
            # Fallback: Add text boxes
            title_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(0.5))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = vuln["description"] or "No title provided"
            p.font.size = Pt(18)
            p.font.bold = True
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            severity_box = slide.shapes.add_textbox(Inches(1), Inches(1.7), Inches(11), Inches(0.5))
            tf = severity_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = f"Severity: {vuln['severity'] or 'Unknown'}"
            p.font.size = Pt(14)
            p.font.color.rgb = gColorsRisk.get(vuln["severity"], gColors['black'])
            if vuln["severity"] == "Low":
                p.font.color.rgb = RGBColor(0, 0, 0)
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            desc_label_box = slide.shapes.add_textbox(Inches(1), Inches(2.4), Inches(11), Inches(0.3))
            tf = desc_label_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = "Description"
            p.font.size = Pt(12)
            p.font.bold = True
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            desc_box = slide.shapes.add_textbox(Inches(1), Inches(2.7), Inches(11), Inches(1.5))
            tf = desc_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = sections.get("Description", "Not provided.")
            p.font.size = Pt(12)
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            impact_label_box = slide.shapes.add_textbox(Inches(1), Inches(4.3), Inches(11), Inches(0.3))
            tf = impact_label_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = "Impact"
            p.font.size = Pt(12)
            p.font.bold = True
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            impact_box = slide.shapes.add_textbox(Inches(1), Inches(4.6), Inches(11), Inches(1.5))
            tf = impact_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = sections.get("Impact", "Not provided.")
            p.font.size = Pt(12)
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            rec_label_box = slide.shapes.add_textbox(Inches(1), Inches(6.2), Inches(11), Inches(0.3))
            tf = rec_label_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = "Recommendation"
            p.font.size = Pt(12)
            p.font.bold = True
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

            rec_box = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11), Inches(1.5))
            tf = rec_box.text_frame
            tf.word_wrap = True
            p = tf.add_paragraph()
            p.text = sections.get("Recommendation", "Not provided.")
            p.font.size = Pt(12)
            p.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT

        if int(vuln.get("attachments", {}).get("total", 0)) > 0:
            print(f"Retrieving attachments for vulnerability {vuln['description']}")
            retrieve_vuln_attachments(presentation, slide_layouts, session, headers, vuln["uuid"])

    _add_slide(presentation, "divider dark blue", slide_layouts).placeholders[0].text = "what's next"

    # FINDINGS DUE DATE slide
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "findings due date overview"
    shape = slide.shapes.add_table(len(sorted_vulns) + 1, 3, Inches(0.8), Inches(1.7), Inches(11), Inches(1))
    table = shape.table
    table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(7), Inches(2), Inches(2)
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(0, 2).text = "finding", "severity", "due date"
    due_dates = {"Critical": "14 days", "High": "30 days", "Medium": "45 days", "Low": "60 days",
                 "Info": "9 months"}
    for i, vuln in enumerate(sorted_vulns):
        table.cell(i + 1, 0).text = vuln['description']
        table.cell(i + 1, 1).text = vuln['severity']
        table.cell(i + 1, 2).text = due_dates.get(vuln['severity'], "N/A")

    # Slide 1 - START
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "steps for an efficient fix validation (retest)"

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.clear()

    p1 = tf.add_paragraph()
    p1.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT
    run1 = p1.add_run()
    run1.text = "Our shared goal is to verify and close security findings as quickly and efficiently as possible. When you mark a vulnerability as \"solved,\" our retest process begins. Here’s how you can help make that process seamless and fast."

    tf.add_paragraph()
    p_heading = tf.add_paragraph()
    run_heading = p_heading.add_run()
    run_heading.text = "1. Retest Readiness"
    run_heading.font.size = Pt(18)

    tf.add_paragraph()
    p2 = tf.add_paragraph()
    run2a = p2.add_run()
    run2a.text = "To validate your fix, "
    run2b = p2.add_run()
    run2b.text = "we need the same access we had during the original test"
    run2b.font.color.rgb = RGBColor(0x98, 0x00, 0x00)
    run2c = p2.add_run()
    run2c.text = ". Delays in access are the #1 reason for delays in closing findings."

    tf.add_paragraph()
    p_b1 = tf.add_paragraph()
    p_b1.bullet = True
    p_b1.level = 0
    run_b1a = p_b1.add_run()
    run_b1a.text = "- Test accounts: "
    run_b1a.font.bold = True
    run_b1b = p_b1.add_run()
    run_b1b.text = "Ensure test accounts remain functional or provide new ones"

    p_b2 = tf.add_paragraph()
    p_b2.bullet = True
    p_b2.level = 0
    run_b2a = p_b2.add_run()
    run_b2a.text = "- Network Access: "
    run_b2a.font.bold = True
    run_b2b = p_b2.add_run()
    run_b2b.text = "If not internet facing, help us with the connectivity whenever possible"

    p_b3 = tf.add_paragraph()
    p_b3.bullet = True
    p_b3.level = 0
    run_b3a = p_b3.add_run()
    run_b3a.text = "- Notify changes: "
    run_b3a.font.bold = True
    run_b3b = p_b3.add_run()
    run_b3b.text = "If the application URL or environment changes, please let us know"

    tf.add_paragraph()
    p_last = tf.add_paragraph()
    run_last = p_last.add_run()
    run_last.text = "A ready environment means we can start validating immediately. A blocked environment means delays for everyone."
    # Slide 1 - END

    # Slide 2 - START
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "steps for an efficient fix validation (retest)"

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.clear()

    p1 = tf.add_paragraph()
    p1.alignment = PP_PARAGRAPH_ALIGNMENT.LEFT
    run1 = p1.add_run()
    run1.text = "Our shared goal is to verify and close security findings as quickly and efficiently as possible. When you mark a vulnerability as \"solved,\" our retest process begins. Here’s how you can help make that process seamless and fast."

    tf.add_paragraph()
    p_heading = tf.add_paragraph()
    run_heading = p_heading.add_run()
    run_heading.text = "2. Provide Clear Fix Context & Evidence"
    run_heading.font.size = Pt(18)

    tf.add_paragraph()
    p2 = tf.add_paragraph()
    run2a = p2.add_run()
    run2a.text = "When you mark a finding as \"solved\" in the platform, "
    run2b = p2.add_run()
    run2b.text = "tell us about the fix. The context is invaluable."
    run2b.font.color.rgb = RGBColor(0x98, 0x00, 0x00)

    tf.add_paragraph()
    p_b1 = tf.add_paragraph()
    p_b1.bullet = True
    p_b1.level = 0
    run_b1a = p_b1.add_run()
    run_b1a.text = "- Add a Comment: "
    run_b1a.font.bold = True
    run_b1b = p_b1.add_run()
    run_b1b.text = "In the vulnerability ticket, briefly describe the change you made."

    p_b2 = tf.add_paragraph()
    p_b2.bullet = True
    p_b2.level = 1
    run_b2a = p_b2.add_run()
    run_b2a.text = "  Example: \"We have implemented server-side validation on the user profile form to sanitize input and prevent XSS.\" "
    run_b2a.font.italic = True
    run_b2a.font.size = Pt(11)

    tf.add_paragraph()
    p_b3 = tf.add_paragraph()
    p_b3.bullet = True
    p_b3.level = 0
    run_b3a = p_b3.add_run()
    run_b3a.text = "- Attach Evidence: "
    run_b3a.font.bold = True
    run_b3b = p_b3.add_run()
    run_b3b.text = "For straightforward changes and fixes, a screenshot can significantly speed up validation."

    p_b4 = tf.add_paragraph()
    p_b4.bullet = True
    p_b4.level = 1
    run_b4a = p_b4.add_run()
    run_b4a.text = "  Good examples: A screenshot showing a security header has been enabled, a directory listing has been disabled, or a debug mode has been turned off. "
    run_b4a.font.italic = True
    run_b4a.font.size = Pt(11)
    # Slide 2 - END

    # Vulnerability Life Cycle and Retest Info Slides
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "vulnerability life cycle"
    slide.shapes.add_picture(vuln_life_cycle_image_stream, Inches(1), Inches(1.5), width=Inches(11), height=Inches(5))

    _add_slide(presentation, "questions", slide_layouts)
    _add_slide(presentation, "divider dark blue", slide_layouts).placeholders[0].text = "appendix"

    # APPENDIX slide
    slide = _add_slide(presentation, "title only", slide_layouts)
    slide.placeholders[0].text = "group of tests performed"
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(0.5))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    run = p.add_run()
    run.text = "example of tests performed during the pentest campaign"
    p.alignment = PP_PARAGRAPH_ALIGNMENT.CENTER
    shape = slide.shapes.add_table(len(group_of_tests), 3, Inches(0.8), Inches(2), Inches(11.5), Inches(1))
    table = shape.table
    table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(0.5), Inches(3), Inches(8)
    for i, (key, (group, example)) in enumerate(group_of_tests.items()):
        table.cell(i, 0).text, table.cell(i, 1).text, table.cell(i, 2).text = key, group, example

    _add_slide(presentation, "last page", slide_layouts)

    # --- Save and Upload ---
    print("Saving presentation to memory...")
    output_stream = BytesIO()
    presentation.save(output_stream)
    output_stream.seek(0)

    output_filename = (
        f"{test['id']} - Restitution Meeting - "
        f"{slugify(asset['name'])} - {datetime.now().strftime('%Y-%m')}.pptx"
    )

    print(f"Uploading '{output_filename}' to Google Drive...")
    file_link = upload_drive_file(drive_service, db_drive_folder_id, output_filename, output_stream)
    print("Upload complete.")

    return {
        "message": "Presentation generated successfully!",
        "driveLink": file_link,
        "warnings": healthy_check
    }